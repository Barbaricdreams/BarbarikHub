"""Read-only extraction for local QA. Requires bundled openpyxl; never writes the workbook."""
import json, pathlib, sys
import openpyxl

source, output = map(pathlib.Path, sys.argv[1:3])
formulas = openpyxl.load_workbook(source, data_only=False)
values = openpyxl.load_workbook(source, data_only=True)
names = ['Draft - AI', 'September 26', 'August 26', 'July 26', 'June 26']
names += [name for name in formulas.sheetnames if name.startswith('May') and '26' in name]
snapshots = []
for name in names:
    s, v = formulas[name], values[name]
    snapshots.append({
        'id': formulas.sheetnames.index(name), 'title': name,
        'values': [[v.cell(r,c).value if v.cell(r,c).value is not None else '' for c in range(1, min(s.max_column,40)+1)] for r in range(1,min(s.max_row,160)+1)],
        'formulas': {c.coordinate:c.value for row in s for c in row if c.data_type=='f'},
        'checkboxes': [c.coordinate for row in v for c in row if isinstance(c.value,bool)],
        'merges': [{'startRow':m.min_row,'endRow':m.max_row,'startCol':m.min_col,'endCol':m.max_col} for m in s.merged_cells.ranges]
    })
output.parent.mkdir(parents=True,exist_ok=True)
output.write_text(json.dumps(snapshots,default=str),encoding='utf-8')
print(f'Read {len(snapshots)} reference tabs. Private snapshots saved locally; source unchanged.')
